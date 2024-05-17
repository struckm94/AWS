import os
import boto3
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.datavalidation import DataValidation

def retrieve_security_groups(region_name=None):
    try:
        region_name = region_name or os.getenv('AWS_REGION', 'us-west-2')
        session = boto3.Session(
            aws_access_key_id=os.getenv('AWS_ACCESS_KEY_ID'),
            aws_secret_access_key=os.getenv('AWS_SECRET_ACCESS_KEY'),
            region_name=region_name
        )
        ec2 = session.resource('ec2')
        return list(ec2.security_groups.all())
    except Exception as e:
        print(f"An error occurred: {e}")
        return []

def add_data_validation(worksheet, max_row):
    # Create data validation for the 'Action' column
    # showDropDown=False needs to be set to "False" to make the dropdown visible - Reference: https://openpyxl.readthedocs.io/en/latest/validation.html
    dv = DataValidation(type="list", formula1='"Keep,Delete,Unknown"', showDropDown=False)
    worksheet.add_data_validation(dv)
    dv.add(f'C3:C{max_row}')  # Apply data validation from row 3 to max_row in column C

def create_excel_file(security_groups):
    workbook = Workbook()
    workbook.remove(workbook.active)  # Remove the default sheet

    for sg in security_groups:
        data = prepare_data(sg)
        df = pd.DataFrame(data)
        if df.empty:
            continue  # Skip empty data frames

        worksheet = workbook.create_sheet(title=sg.id[:31])  # Create new sheet with the SG ID
        # Add header for security group name
        worksheet.append([f"Security Group Name: {sg.group_name}"])
        # Merge cells for the header
        worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
        
        # Add data from DataFrame to the worksheet
        for r in dataframe_to_rows(df, index=False, header=True):
            worksheet.append(r)

        # Set data validation for 'Action' column
        add_data_validation(worksheet, len(df) + 2)

    workbook.save('Security_Groups.xlsx')

def prepare_data(sg):
    data = {
        'Rule Name': [],
        'IP': [],
        'Action (Keep/Delete/Unknown)': [],
        'Name': [],
        'Description': []
    }
    for rule in sg.ip_permissions:
        for ip_range in rule.get('IpRanges', []):
            rule_name = f"PORT {rule.get('FromPort', 'N/A')}" if rule.get('FromPort') else "All Ports"
            data['Rule Name'].append(rule_name)
            data['IP'].append(ip_range['CidrIp'])
            data['Action (Keep/Delete/Unknown)'].append('')
            data['Name'].append('')
            data['Description'].append(rule.get('Description', 'No description'))
    return data

if __name__ == "__main__":
    security_groups = retrieve_security_groups()
    if security_groups:
        create_excel_file(security_groups)
    else:
        print("No security groups found or an error occurred.")
